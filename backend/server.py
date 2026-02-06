from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional
import uuid
from datetime import datetime, timezone
import io
import re
import PyPDF2
import docx
import markdown
from bs4 import BeautifulSoup
import requests
from openai import OpenAI
from google import genai
from google.genai import types

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client_db = AsyncIOMotorClient(mongo_url)
db = client_db[os.environ['DB_NAME']]

openai_client = OpenAI(api_key=os.environ.get('OPENAI_API_KEY'))
gemini_client = genai.Client(api_key=os.environ.get('GEMINI_API_KEY'))

app = FastAPI()
api_router = APIRouter(prefix="/api")

class Document(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    title: str
    content: str
    sentences: List[str]
    source_type: str
    source_url: Optional[str] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class DocumentCreate(BaseModel):
    title: str
    content: str
    source_type: str
    source_url: Optional[str] = None

class URLInput(BaseModel):
    url: str

class ChatMessage(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    document_id: str
    role: str
    content: str
    model: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class ChatRequest(BaseModel):
    document_id: str
    message: str
    model: str
    context: Optional[str] = None

class TTSRequest(BaseModel):
    text: str
    voice: Optional[str] = "alloy"

def extract_sentences(text: str) -> List[str]:
    text = re.sub(r'\s+', ' ', text).strip()
    sentences = re.split(r'(?<=[.!?])\s+(?=[A-Z])', text)
    return [s.strip() for s in sentences if s.strip()]

def extract_text_from_pdf(file_bytes: bytes) -> str:
    pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text() + "\n"
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join([para.text for para in doc.paragraphs])

def extract_text_from_markdown(file_bytes: bytes) -> str:
    md_text = file_bytes.decode('utf-8')
    html = markdown.markdown(md_text)
    return BeautifulSoup(html, 'html.parser').get_text()

def scrape_url_content(url: str) -> str:
    try:
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, 'html.parser')
        for script in soup(["script", "style"]):
            script.decompose()
        return soup.get_text(separator='\n', strip=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to scrape URL: {str(e)}")

@api_router.post("/documents/upload", response_model=Document)
async def upload_document(file: UploadFile = File(...)):
    try:
        file_bytes = await file.read()
        filename = file.filename.lower()
        
        if filename.endswith('.pdf'):
            content = extract_text_from_pdf(file_bytes)
        elif filename.endswith('.docx'):
            content = extract_text_from_docx(file_bytes)
        elif filename.endswith(('.md', '.markdown')):
            content = extract_text_from_markdown(file_bytes)
        elif filename.endswith('.txt'):
            content = file_bytes.decode('utf-8')
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format")
        
        sentences = extract_sentences(content)
        doc_data = {
            "title": file.filename,
            "content": content,
            "sentences": sentences,
            "source_type": "file",
            "source_url": None
        }
        doc = Document(**doc_data)
        doc_dict = doc.model_dump()
        doc_dict['created_at'] = doc_dict['created_at'].isoformat()
        await db.documents.insert_one(doc_dict)
        return doc
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@api_router.post("/documents/url", response_model=Document)
async def add_url_document(input_data: URLInput):
    try:
        content = scrape_url_content(input_data.url)
        sentences = extract_sentences(content)
        title = input_data.url.split('//')[-1].split('/')[0]
        
        doc_data = {
            "title": title,
            "content": content,
            "sentences": sentences,
            "source_type": "url",
            "source_url": input_data.url
        }
        doc = Document(**doc_data)
        doc_dict = doc.model_dump()
        doc_dict['created_at'] = doc_dict['created_at'].isoformat()
        await db.documents.insert_one(doc_dict)
        return doc
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@api_router.get("/documents", response_model=List[Document])
async def get_documents():
    docs = await db.documents.find({}, {"_id": 0}).sort("created_at", -1).to_list(100)
    for doc in docs:
        if isinstance(doc['created_at'], str):
            doc['created_at'] = datetime.fromisoformat(doc['created_at'])
    return docs

@api_router.get("/documents/{doc_id}", response_model=Document)
async def get_document(doc_id: str):
    doc = await db.documents.find_one({"id": doc_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    if isinstance(doc['created_at'], str):
        doc['created_at'] = datetime.fromisoformat(doc['created_at'])
    return doc

@api_router.delete("/documents/{doc_id}")
async def delete_document(doc_id: str):
    result = await db.documents.delete_one({"id": doc_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Document not found")
    await db.chat_messages.delete_many({"document_id": doc_id})
    return {"message": "Document deleted"}

@api_router.post("/tts")
async def text_to_speech(request: TTSRequest):
    try:
        response = openai_client.audio.speech.create(
            model="tts-1",
            voice=request.voice,
            input=request.text[:4096]
        )
        audio_bytes = io.BytesIO()
        for chunk in response.iter_bytes():
            audio_bytes.write(chunk)
        audio_bytes.seek(0)
        return StreamingResponse(audio_bytes, media_type="audio/mpeg")
    except Exception as e:
        logger.error(f"TTS error: {str(e)}")
        if "429" in str(e) or "rate_limit" in str(e).lower():
            raise HTTPException(status_code=429, detail="OpenAI API rate limit exceeded. Please try browser TTS or wait a few minutes.")
        raise HTTPException(status_code=500, detail=f"TTS failed: {str(e)}")

@api_router.post("/chat", response_model=ChatMessage)
async def chat_with_ai(request: ChatRequest):
    try:
        doc = await db.documents.find_one({"id": request.document_id}, {"_id": 0})
        if not doc:
            raise HTTPException(status_code=404, detail="Document not found")
        
        context = request.context or doc['content'][:8000]
        system_prompt = f"You are a helpful AI assistant. Answer questions based on this documentation:\n\n{context}"
        
        if request.model.startswith('gpt'):
            try:
                response = openai_client.chat.completions.create(
                    model="gpt-4o",
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": request.message}
                    ],
                    timeout=30
                )
                answer = response.choices[0].message.content
            except Exception as e:
                logger.error(f"OpenAI chat error: {str(e)}")
                if "429" in str(e) or "rate_limit" in str(e).lower():
                    raise HTTPException(status_code=429, detail="OpenAI API rate limit exceeded. Please try Gemini model or wait a few minutes.")
                raise
        else:
            try:
                prompt = f"{system_prompt}\n\nUser: {request.message}"
                response = gemini_client.models.generate_content(
                    model='gemini-2.0-flash-exp',
                    contents=prompt
                )
                answer = response.text
            except Exception as e:
                logger.error(f"Gemini chat error: {str(e)}")
                raise HTTPException(status_code=500, detail=f"Gemini API error: {str(e)}")
        
        msg_data = {
            "document_id": request.document_id,
            "role": "assistant",
            "content": answer,
            "model": request.model
        }
        msg = ChatMessage(**msg_data)
        msg_dict = msg.model_dump()
        msg_dict['created_at'] = msg_dict['created_at'].isoformat()
        await db.chat_messages.insert_one(msg_dict)
        return msg
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Chat error: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@api_router.get("/chat/{document_id}", response_model=List[ChatMessage])
async def get_chat_history(document_id: str):
    messages = await db.chat_messages.find(
        {"document_id": document_id}, {"_id": 0}
    ).sort("created_at", 1).to_list(100)
    for msg in messages:
        if isinstance(msg['created_at'], str):
            msg['created_at'] = datetime.fromisoformat(msg['created_at'])
    return messages

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client_db.close()